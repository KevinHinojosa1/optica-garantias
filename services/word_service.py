from datetime import datetime
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from models.historial import HistorialConsulta
from models.cliente import Cliente
from services.tiendas_service import TiendasService
import os
from config import settings, LOGO_OFICIAL_PATH

class WordService:
    @staticmethod
    def _set_cell_background(cell, color_hex):
        """Set background color for a table cell"""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), color_hex)
        tcPr.append(shading)

    @classmethod
    def generar_informe_consulta(cls, historial: HistorialConsulta, cliente: Cliente | None) -> bytes:
        doc = Document()
        
        # Margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)
            
        teal_color = RGBColor(13, 148, 136) # #0D9488
        light_teal = "E6F7F5"
        light_gray = "F3F4F6"
        
        # 1. Header
        header_table = doc.add_table(rows=1, cols=2)
        # Give column 0 most of the width, column 1 the rest
        header_table.columns[0].width = Inches(4.5)
        header_table.columns[1].width = Inches(2.0)
        
        # Title in the left cell
        title_cell = header_table.rows[0].cells[0]
        title_para = title_cell.paragraphs[0]
        title_run = title_para.add_run("INFORME DE REVISIÓN DE LENTES")
        title_run.bold = True
        title_run.font.size = Pt(16)
        
        # Subtitle
        sub_run = title_para.add_run("\nExplicación clara para usted")
        sub_run.font.size = Pt(10)
        sub_run.font.color.rgb = RGBColor(128, 128, 128)
        
        # Logo in the right cell
        logo_cell = header_table.rows[0].cells[1]
        logo_para = logo_cell.paragraphs[0]
        logo_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if os.path.exists(LOGO_OFICIAL_PATH):
            logo_para.add_run().add_picture(LOGO_OFICIAL_PATH, width=Inches(1.8))
            
        doc.add_paragraph() # Spacer
        
        # Date formatting
        meses = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
        fecha_actual = datetime.now()
        fecha_str = f"{fecha_actual.day} de {meses[fecha_actual.month-1]} de {fecha_actual.year}"

        # Info Fallbacks
        orden_trabajo = str(cliente.numero_factura) if cliente and cliente.numero_factura else "N/A"
        local_name = "N/A"
        if cliente and cliente.tienda:
            tienda_info = TiendasService.resolver_para_cliente(cliente.tienda)
            local_name = tienda_info.get("nombre", cliente.tienda)
            
        producto = cliente.producto if cliente and cliente.producto else "Lentes de medida"

        # 2. Info Table
        table = doc.add_table(rows=2, cols=4)
        table.style = 'Table Grid'
        
        cells = table.rows[0].cells
        cells[0].text = "Orden de trabajo"
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].paragraphs[0].runs[0].font.color.rgb = teal_color
        cls._set_cell_background(cells[0], light_teal)
        
        cells[1].text = orden_trabajo
        
        cells[2].text = "Local"
        cells[2].paragraphs[0].runs[0].bold = True
        cells[2].paragraphs[0].runs[0].font.color.rgb = teal_color
        cls._set_cell_background(cells[2], light_teal)
        
        cells[3].text = local_name
        
        cells = table.rows[1].cells
        cells[0].text = "Fecha"
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].paragraphs[0].runs[0].font.color.rgb = teal_color
        cls._set_cell_background(cells[0], light_teal)
        
        cells[1].text = fecha_str
        
        cells[2].text = "Producto"
        cells[2].paragraphs[0].runs[0].bold = True
        cells[2].paragraphs[0].runs[0].font.color.rgb = teal_color
        cls._set_cell_background(cells[2], light_teal)
        
        cells[3].text = producto
        
        doc.add_paragraph() # Spacer

        # 3. Section: ¿Qué revisamos?
        p = doc.add_paragraph()
        run = p.add_run("¿Qué revisamos?")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = teal_color
        
        doc.add_paragraph("Revisamos cuidadosamente sus lentes por ambos lados, sus bordes y las zonas que tocan la montura. El objetivo fue conocer qué ocurrió con la superficie y verificar si el problema podía haberse originado durante la fabricación.")

        # 4. Section: ¿Qué encontramos?
        p = doc.add_paragraph()
        run = p.add_run("¿Qué encontramos?")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = teal_color
        
        que_observamos_text = historial.que_observamos or historial.motivo or "Se observan alteraciones en la superficie del lente."
        que_causa_text = historial.que_causa or (
            "Este tipo de cambio puede estar relacionado con un defecto en el proceso de fabricación o en los materiales utilizados."
            if historial.veredicto == "APLICA"
            else "Este tipo de daño puede aparecer durante el uso por situaciones como calor intenso, cambios bruscos de temperatura, presión, roce frecuente o productos de limpieza no adecuados para lentes."
        )
        problema_fabricacion_text = historial.problema_fabricacion or (
            "Sí, se detectaron señales de un defecto de fabricación."
            if historial.veredicto == "APLICA"
            else ("No es posible determinar con la información actual. Se requiere una revisión presencial." if historial.veredicto == "IMAGEN NO CLARA" else "No, no encontramos señales de que el problema se haya originado durante la fabricación.")
        )

        table_enc = doc.add_table(rows=3, cols=2)
        table_enc.style = 'Table Grid'
        
        rows = [
            ("¿Qué observamos?", que_observamos_text),
            ("¿Qué puede causar este tipo de cambio?", que_causa_text),
            ("¿Encontramos un problema de fabricación?", problema_fabricacion_text)
        ]
        
        for i, (q, a) in enumerate(rows):
            cells = table_enc.rows[i].cells
            cells[0].text = q
            cells[0].paragraphs[0].runs[0].bold = True
            cells[0].paragraphs[0].runs[0].font.color.rgb = teal_color
            cls._set_cell_background(cells[0], light_teal)
            cells[1].text = a

        doc.add_paragraph()

        # 5. Section: ¿Cuál es el resultado de la revisión?
        p = doc.add_paragraph()
        run = p.add_run("¿Cuál es el resultado de la revisión?")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = teal_color
        
        resultado_revision_text = historial.resultado_revision or (
            f"Después de revisar sus lentes, encontramos que el cambio visible puede estar relacionado con situaciones que ocurren durante el uso o con el proceso de fabricación. {historial.fundamento or ''}"
        )
        
        res_table = doc.add_table(rows=1, cols=1)
        res_cell = res_table.rows[0].cells[0]
        cls._set_cell_background(res_cell, light_gray)
        res_run = res_cell.paragraphs[0].add_run(resultado_revision_text)
        res_run.bold = True
        
        doc.add_paragraph()

        # 6. Warranty Decision Line
        if historial.veredicto == "APLICA":
            decision_text = "En esta ocasión la garantía sí aplica, porque en la revisión encontramos señales de que el problema se originó durante la fabricación."
        elif historial.veredicto == "NO APLICA":
            decision_text = "En esta ocasión no aplica, porque en la revisión no encontramos señales de que el problema se haya originado durante la fabricación."
        else:
            decision_text = "No es posible determinar con la información actual. Se requiere una revisión presencial."
            
        p = doc.add_paragraph()
        p.add_run("¿La garantía aplica en este caso? ").bold = True
        p.add_run(decision_text)

        # 7. Section: ¿Cómo puede cuidar mejor sus lentes?
        p = doc.add_paragraph()
        run = p.add_run("¿Cómo puede cuidar mejor sus lentes?")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = teal_color
        
        bullets = [
            ("Para limpiarlos: ", "use un paño de microfibra limpio y un producto adecuado para lentes. Evite limpiadores del hogar o productos químicos fuertes."),
            ("Para protegerlos del calor: ", "no los deje dentro del vehículo expuestos al sol ni cerca de fuentes de calor."),
            ("Evite cambios bruscos de temperatura: ", "por ejemplo, no exponga los lentes inmediatamente a calor intenso después de haber estado en un ambiente muy frío."),
            ("Para evitar presión y roce: ", "no los coloque con los lentes hacia abajo ni debajo de objetos pesados."),
            ("Cuando no los use: ", "guárdelos en su estuche para protegerlos.")
        ]
        
        for bold_part, text_part in bullets:
            bp = doc.add_paragraph(style='List Bullet')
            bp.add_run(bold_part).bold = True
            bp.add_run(text_part)

        doc.add_paragraph()
        
        # 8. Commitment Block
        comp_table = doc.add_table(rows=1, cols=1)
        comp_cell = comp_table.rows[0].cells[0]
        cls._set_cell_background(comp_cell, light_teal)
        
        comp_p = comp_cell.paragraphs[0]
        comp_p.add_run("Nuestro compromiso con usted: ").bold = True
        comp_p.runs[0].font.color.rgb = teal_color
        comp_text = comp_p.add_run("queremos que conozca de forma clara qué observamos, por qué llegamos a este resultado y cómo puede cuidar mejor sus lentes. Si tiene alguna duda sobre esta revisión, nuestro equipo de Servicio al Cliente está disponible para orientarlo.")
        comp_text.font.color.rgb = teal_color
        
        doc.add_paragraph()
        doc.add_paragraph()

        # 9. Signature Block
        sig_table = doc.add_table(rows=1, cols=2)
        sig_left = sig_table.rows[0].cells[0]
        p_left = sig_left.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_left.add_run("________________________\nSERVICIO AL CLIENTE / POSVENTA\nÓPTICA LOS ANDES S.A.").bold = True
        for r in p_left.runs: r.font.color.rgb = teal_color
        
        sig_right = sig_table.rows[0].cells[1]
        p_right = sig_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_right.add_run("________________________\nRESPONSABLE DE REVISIÓN / AUTORIZACIÓN\nÓPTICA LOS ANDES S.A.").bold = True
        for r in p_right.runs: r.font.color.rgb = teal_color

        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
